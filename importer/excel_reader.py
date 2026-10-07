from __future__ import annotations

import hashlib
from pathlib import Path

from openpyxl import load_workbook

from .models import ParsedOrder


class UnsupportedOrderTemplate(ValueError):
    pass


EXPECTED_SHEETS = {"HOJA PEDIDO", "OBRAS", "MATERIALES"}


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_order(path: str | Path) -> ParsedOrder:
    """Lee un pedido Excel.

    La extracción exacta de celdas de la plantilla real se implementará en el
    siguiente hito, tras fijar formalmente el mapa de campos de la muestra
    2024-2026.
    """
    path = Path(path)
    workbook = load_workbook(path, data_only=False, read_only=True)

    if not EXPECTED_SHEETS.issubset(set(workbook.sheetnames)):
        raise UnsupportedOrderTemplate(
            f"Plantilla no reconocida. Hojas encontradas: {workbook.sheetnames}"
        )

    return ParsedOrder(
        source_filename=path.name,
        source_sha256=sha256_file(path),
    )
