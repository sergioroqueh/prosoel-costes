from datetime import datetime

from openpyxl import Workbook

from importer.excel_reader import read_order


def test_reads_legacy_two_discount_template(tmp_path) -> None:
    path = tmp_path / "pedido_legacy.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "HOJA PEDIDO"

    ws["A11"] = "FECHA:"
    ws["C11"] = datetime(2024, 2, 19)
    ws["A12"] = "RESPONSABLE OFICINA"
    ws["C12"] = "ISRAEL"
    ws["A14"] = "ALMACÉN:"
    ws["C14"] = "GRUPO RIAS"
    ws["A18"] = "Nº PEDIDO: 24/"
    ws["C18"] = "229"
    ws["A19"] = "IMPORTE PEDIDO:"
    ws["C19"] = 51.4008
    ws["A20"] = "REFERENCIA OBRA:"
    ws["C20"] = "LOFTS VILLAVERDE"

    headers = [
        "UDS.",
        "REFERENCIA",
        "MATERIAL",
        "PRECIO",
        "DTO. 1",
        "DTO. 2",
        "PRECIO NETO",
        "PRECIO TOTAL",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(24, col).value = value

    values = [
        20,
        "20000472-090",
        "Base de enchufe SIMON 270",
        5.9,
        0.56,
        0.01,
        2.57004,
        51.4008,
    ]
    for col, value in enumerate(values, 1):
        ws.cell(25, col).value = value

    wb.save(path)

    order = read_order(path)

    assert order.order_reference == "24/229"
    assert order.project == "LOFTS VILLAVERDE"
    assert order.template_variant == "legacy_two_discount"
    assert len(order.lines) == 1
    assert order.lines[0].net_unit_price == 2.57004
    assert order.lines[0].total_price == 51.4008
    assert order.lines[0].discount_components_raw == ["0.56", "0.01"]
